"""Opt-in real local DOCX/PPTX rendering and Apple Vision/editppt smoke tests."""
import json
import os
import shutil
import subprocess
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

from test_regressions import LocalCase, SCRIPTS, op, ps, rt


@unittest.skipUnless(os.environ.get("PPT_GEN_RUN_INTEGRATION") == "1", "set PPT_GEN_RUN_INTEGRATION=1 for real local tool checks")
class IntegrationTests(LocalCase):
    def keep(self, path, name):
        directory = os.environ.get("PPT_GEN_TEST_ARTIFACTS_DIR")
        if directory:
            target = Path(directory)
            target.mkdir(parents=True, exist_ok=True)
            shutil.copy2(path, target / name)

    def test_existing_docx_qa_only_can_complete(self):
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Pt
        document = Document()
        style = document.styles["Normal"]
        style.font.name = "PingFang SC"
        style.font.size = Pt(12)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        document.add_heading("Ppt Gen 本地质检测试", 0)
        statement = "本报告仅用于验证已有成品的独立质检流程。"
        document.add_heading("测试范围", 1)
        document.add_paragraph(statement)
        document.add_paragraph("检查中文字体、正文排版、渲染结果，以及输入文件是否保持不变。")
        document.add_heading("资料来源", 1)
        document.add_paragraph("来源：本地人工构造的测试材料，不包含外部研究结论。")
        report = self.root / "existing-report.docx"
        document.save(report)
        report_digest = rt.sha256(report)
        source = self.write("source.txt", statement)
        source_manifest = self.write("source-manifest.json", {"schema_version": 1,
            "source_policy": "faithful", "data_handling": "standard", "ocr_policy": "offline",
            "external_ocr_allowed": False, "ask_for_ocr_token": False, "primary_source_id": "src-1",
            "sources": [{"id": "src-1", **op.binding(source)}]})
        claims = self.write("claim-ledger.json", {"schema_version": 1, "source_cutoff": "2026-09-11",
            "claims": [{"id": "claim-1", "statement": statement, "kind": "user-provided",
                "verification": "user-authority", "source_ids": ["src-1"], "source_locator": "local fixture",
                "used_in_report": True, "report_evidence": [statement], "used_on_slides": []}]})
        cache = self.root / "cache"
        rt.configure_cache(cache)
        with patch.object(rt.subprocess, "run", wraps=subprocess.run) as commands:
            pages = rt.render_pages(report)
            self.assertTrue(pages)
            self.assertEqual(pages, rt.render_pages(report))
            conversions = [call for call in commands.call_args_list if "--convert-to" in call.args[0]]
            self.assertEqual(len(conversions), 1)
        evidence = self.write("render-evidence.json", {"schema_version": 1, "tool": "ppt-gen.render-evidence",
            "passed": True, "artifact": op.binding(report), "renderer": {"name": "libreoffice",
                "version": "local-smoke", "input_sha256": report_digest}, "artifact_page_count": len(pages),
            "pages": [{"page": i, **op.binding(path)} for i, path in enumerate(pages, 1)]})
        project = self.root / "qa-project"
        self.state_command("init", project, "--source-policy", "faithful", "--start-at", "qa-package", "--stop-after", "qa-package", "--deliverable", "qa-report")
        self.state_command("add-input", project, "existing=" + str(report), "--artifact-role", "report", "--stage", "qa-package")
        for name, path in (("sources", source_manifest), ("claims", claims), ("render", evidence)):
            self.state_command("add-input", project, name + "=" + str(path))
        self.state_command("verify", project)
        qa_path = self.root / "qa-report.json"
        command = [sys.executable, "-B", str(SCRIPTS / "validate_delivery.py"), "--report", str(report),
                   "--source-manifest", str(source_manifest), "--claim-ledger", str(claims),
                   "--report-render-evidence", str(evidence), "--project-state", str(project / ps.MANIFEST_NAME),
                   "--cache-dir", str(cache), "--json-out", str(qa_path)]
        checked = subprocess.run(command, capture_output=True, text=True)
        self.assertEqual(checked.returncode, 0, checked.stderr + checked.stdout)
        qa = json.loads(qa_path.read_text())
        self.state_command("set-stage", project, "qa-package", "--status", "completed", "--qa", qa["status"], "--deliverable", "qa-report=" + str(qa_path))
        state = json.loads((project / ps.MANIFEST_NAME).read_text())
        self.assertEqual(state["run_status"], "completed")
        self.assertEqual(state["stages"]["report"]["status"], "skipped")
        self.assertEqual(rt.sha256(report), report_digest)
        self.keep(pages[0], "word-qa-smoke.png")
        self.keep(qa_path, "word-qa-report.json")
        self.keep(report, "word-qa-smoke.docx")

    def test_guarded_prepare_uses_real_offline_recognition(self):
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new("RGB", (1920, 1080), "#f4f6f8")
        draw = ImageDraw.Draw(image)
        font = ImageFont.truetype("/System/Library/Fonts/STHeiti Medium.ttc", 80)
        draw.text((160, 180), "Ppt Gen 离线识别测试", font=font, fill="#182433")
        draw.text((160, 380), "本页不包含外部研究结论", font=font, fill="#182433")
        source = self.root / "01-slide.png"
        image.save(source)
        data = {"schema_version": 1, "ocr_policy": "offline", "ask_for_ocr_token": False,
                "data_handling": "standard", "image_backend": "builtin-allowed", "source_policy": "faithful",
                "source_images": [op.binding(source)], "hashes": {}}
        for key in ("source_manifest", "slide_copy_ledger", "slide_manifest"):
            path = self.write(key + ".json", {"schema_version": 1, "source_policy": "faithful", "data_handling": "standard"})
            data[key] = str(path.resolve())
            data["hashes"][key] = rt.sha256(path)
        handoff = self.write("handoff.json", data)
        execution = self.root / "offline-preparation.json"
        rt.configure_cache(self.root / "ocr-cache")
        op.prepare(handoff, self.root / "run", execution)
        self.assertEqual(op.validate_report(execution, handoff), [])
        tampered = json.loads(execution.read_text())
        tampered["prepare_command"].remove("--no-text-hints")
        tampered_path = self.write("unguarded-preparation.json", tampered)
        self.assertTrue(op.validate_report(tampered_path, handoff))
        transcript = json.loads((self.root / "run/pages/page_001/offline-recognition.json").read_text())
        self.assertIn("离线", "".join(transcript["lines"]))
        self.assertTrue(transcript["blocks"])
        with patch.object(rt.subprocess, "run", wraps=subprocess.run) as commands:
            second = rt.recognize_image(self.root / "run/pages/page_001/source.png")
            self.assertEqual(second["lines"], transcript["lines"])
            self.assertFalse(any(str(call.args[0][0]).endswith("/offline-ocr") for call in commands.call_args_list))
        self.keep(source, "offline-ocr-smoke.png")
        self.keep(execution, "offline-preparation.json")
        self.keep(self.root / "run/pages/page_001/offline-recognition.json", "offline-recognition.json")

    def test_native_pptx_renders_locally(self):
        from pptx import Presentation
        from pptx.util import Inches, Pt
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(16), Inches(9)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1.2), Inches(14), Inches(2))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = "Ppt Gen 可编辑页面测试"
        paragraph.font.name, paragraph.font.size = "PingFang SC", Pt(42)
        body = slide.shapes.add_textbox(Inches(1), Inches(3.5), Inches(13), Inches(2))
        body.text_frame.text = "此处为原生文本框，用于验证本地 PPTX 渲染。"
        body.text_frame.paragraphs[0].font.name = "PingFang SC"
        body.text_frame.paragraphs[0].font.size = Pt(26)
        path = self.root / "native-smoke.pptx"
        deck.save(path)
        rt.configure_cache(self.root / "ppt-cache")
        pages = rt.render_pages(path)
        self.assertEqual(len(pages), 1)
        from PIL import Image
        with Image.open(pages[0]) as rendered:
            self.assertEqual(rendered.width * 9, rendered.height * 16)
        self.keep(pages[0], "ppt-qa-smoke.png")
        self.keep(path, "ppt-qa-smoke.pptx")


if __name__ == "__main__":
    unittest.main()
